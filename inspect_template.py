import openpyxl

template_path = r"data/Metadata Template(1).xlsx"

wb = openpyxl.load_workbook(template_path, data_only=False)

print("Sheets:")
for sheet in wb.sheetnames:
    print(f" - {sheet}")

print("\n--- SEARCHING FOR ATTRIBUTE SECTION ---")

for ws in wb.worksheets:
    print(f"\nSHEET: {ws.title}")

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                value = str(cell.value).strip()

                if (
                    "Attribute #1" in value
                    or "Attribute Label" in value
                    or "Attribute Definition" in value
                    or "Range Domain Minimum" in value
                    or "Range Domain Maximum" in value
                    or "Attribute Unit" in value
                    or "Attribute #2" in value
                ):
                    print(
                        f"Cell: {cell.coordinate:<8} "
                        f"Row: {cell.row:<4} "
                        f"Column: {cell.column:<3} "
                        f"Value: {cell.value}"
                    )