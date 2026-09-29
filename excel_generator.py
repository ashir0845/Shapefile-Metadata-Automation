import os
import openpyxl
from copy import copy

from metadata_builder import build_metadata


DBF_PATH = r"data/BLOCK_2024/BLOCK__INDIA_2024.dbf"
MASTER_PATH = r"data/Metadata Template(1).xlsx"

TEMPLATE_PATH = r"data/Metadata Template(1).xlsx"
OUTPUT_PATH = r"output/Generated_Metadata.xlsx"

SHEET_NAME = "Meatadata_State_2001"

ATTRIBUTE_START_ROW = 40
ATTRIBUTE_BLOCK_SIZE = 7

# Row 61 is a merged row in the original template.
ATTRIBUTE_4_START_ROW = 62


def copy_row_format(ws, source_row, target_row):

    if source_row in ws.row_dimensions:
        ws.row_dimensions[target_row].height = (
            ws.row_dimensions[source_row].height
        )

    for column in range(1, 4):

        source_cell = ws.cell(source_row, column)
        target_cell = ws.cell(target_row, column)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)

        target_cell.number_format = source_cell.number_format
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.font = copy(source_cell.font)


def copy_attribute_block(ws, source_start_row, target_start_row):

    for offset in range(ATTRIBUTE_BLOCK_SIZE):

        source_row = source_start_row + offset
        target_row = target_start_row + offset

        copy_row_format(
            ws,
            source_row,
            target_row
        )


def write_attribute_block(
    ws,
    start_row,
    index,
    metadata
):

    label_row = start_row + 1
    definition_row = start_row + 2
    definition_source_row = start_row + 3
    min_row = start_row + 4
    max_row = start_row + 5
    unit_row = start_row + 6

    # --------------------------------
    # Attribute number
    # --------------------------------

    ws.cell(start_row, 1).value = (
        f"Attribute #{index}"
    )

    # --------------------------------
    # Attribute Label
    # --------------------------------

    ws.cell(
        label_row,
        1
    ).value = "Attribute Label"

    ws.cell(
        label_row,
        2
    ).value = metadata.get("name", "")

    ws.cell(
        label_row,
        3
    ).value = "From Shape File"

    # --------------------------------
    # Attribute Definition
    # --------------------------------

    ws.cell(
        definition_row,
        1
    ).value = "Attribute Definition"

    ws.cell(
        definition_row,
        2
    ).value = (
        metadata.get("definition")
        or ""
    )

    ws.cell(
        definition_row,
        3
    ).value = (
        "From Master"
        if metadata.get("definition")
        else ""
    )

    # --------------------------------
    # Attribute Definition Source
    # --------------------------------

    ws.cell(
        definition_source_row,
        1
    ).value = "Attribute Definition Source"

    ws.cell(
        definition_source_row,
        2
    ).value = (
        metadata.get("definition_source")
        or ""
    )

    ws.cell(
        definition_source_row,
        3
    ).value = (
        "From Master"
        if metadata.get("definition_source")
        else ""
    )

    # --------------------------------
    # Minimum
    # --------------------------------

    ws.cell(
        min_row,
        1
    ).value = "Range Domain Minimum"

    minimum = metadata.get("min")

    ws.cell(
        min_row,
        2
    ).value = (
        minimum
        if minimum is not None
        else ""
    )

    ws.cell(
        min_row,
        3
    ).value = (
        "From Shape File"
        if minimum is not None
        else ""
    )

    # --------------------------------
    # Maximum
    # --------------------------------

    ws.cell(
        max_row,
        1
    ).value = "Range Domain Maximum"

    maximum = metadata.get("max")

    ws.cell(
        max_row,
        2
    ).value = (
        maximum
        if maximum is not None
        else ""
    )

    ws.cell(
        max_row,
        3
    ).value = (
        "From Shape File"
        if maximum is not None
        else ""
    )

    # --------------------------------
    # Unit
    # --------------------------------

    ws.cell(
        unit_row,
        1
    ).value = "Attribute Unit of Measurement"

    unit = metadata.get("unit")

    ws.cell(
        unit_row,
        2
    ).value = (
        unit
        if unit is not None
        else ""
    )

    ws.cell(
        unit_row,
        3
    ).value = (
        "From Shape File"
        if unit
        else ""
    )


def create_excel():

    print("Loading Excel template...")

    workbook = openpyxl.load_workbook(
        TEMPLATE_PATH,
        data_only=False
    )

    ws = workbook[SHEET_NAME]

    print("Template loaded successfully!")

    # --------------------------------
    # Build metadata
    # --------------------------------

    print("Building metadata...")

    metadata = build_metadata(
        DBF_PATH,
        MASTER_PATH
    )

    metadata_list = metadata["fields"]

    print(
        f"Metadata fields received: "
        f"{len(metadata_list)}"
    )

    print(
        f"Total records: "
        f"{metadata['total_records']}"
    )

    # --------------------------------
    # Clear existing sample attributes
    # --------------------------------

    for row in range(40, 61):

        for column in range(1, 4):

            ws.cell(
                row,
                column
            ).value = None

    # --------------------------------
    # Generate attributes
    # --------------------------------

    for index, field_metadata in enumerate(
        metadata_list,
        start=1
    ):

        # Existing blocks
        if index <= 3:

            start_row = (
                ATTRIBUTE_START_ROW
                + (index - 1)
                * ATTRIBUTE_BLOCK_SIZE
            )

        # New blocks
        else:

            start_row = (
                ATTRIBUTE_4_START_ROW
                + (index - 4)
                * ATTRIBUTE_BLOCK_SIZE
            )

            # Copy formatting from Attribute #3
            copy_attribute_block(
                ws,
                54,
                start_row
            )

        write_attribute_block(
            ws,
            start_row,
            index,
            field_metadata
        )

    # --------------------------------
    # Create output directory
    # --------------------------------

    os.makedirs(
        os.path.dirname(OUTPUT_PATH),
        exist_ok=True
    )

    # --------------------------------
    # Save generated workbook
    # --------------------------------

    workbook.save(
        OUTPUT_PATH
    )

    print()
    print("================================")
    print("Excel generated successfully!")
    print("================================")
    print(
        f"Total attributes: "
        f"{len(metadata_list)}"
    )
    print(
        f"Total records: "
        f"{metadata['total_records']}"
    )
    print(
        f"Output: {OUTPUT_PATH}"
    )


if __name__ == "__main__":
    create_excel()