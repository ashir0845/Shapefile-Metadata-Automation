import os
import shutil
import tempfile
import zipfile
import copy
import uuid

from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    UploadFile,
)

from fastapi.responses import FileResponse

from backend.services.metadata_service import (
    build_metadata,
)
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
from sqlalchemy.orm import Session

from backend.database import get_db
from backend.models.history import GeneratedFileHistory

from backend.auth_utils import get_current_user
from backend.models.user import User

# =========================================================
# ROUTER
# =========================================================

router = APIRouter(
    prefix="/metadata",
    tags=["Metadata"],
)


# =========================================================
# PATHS
# =========================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)

TEMPLATE_DIR = os.path.join(
    BASE_DIR,
    "templates",
)

TEMPLATE_PATH = os.path.join(
    TEMPLATE_DIR,
    "Metadata Template(1).xlsx",
)


# =========================================================
# TEMP DIRECTORY
# =========================================================

UPLOAD_ROOT = os.path.join(
    tempfile.gettempdir(),
    "gis_metadata_generator",
)

os.makedirs(
    UPLOAD_ROOT,
    exist_ok=True,
)

# Generated Excel files are kept here (inside the backend folder)
# so the OS cannot clear them and History downloads keep working.
GENERATED_DIR = os.path.join(
    BASE_DIR,
    "generated",
)

os.makedirs(
    GENERATED_DIR,
    exist_ok=True,
)


# =========================================================
# FILE HELPERS
# =========================================================


def find_shapefile_components(directory: str) -> dict:
    """Find required shapefile components recursively."""
    required = {".shp", ".shx", ".dbf", ".prj"}
    found = {}

    for root, dirs, files in os.walk(directory):
        dirs[:] = [
            folder
            for folder in dirs
            if folder != "__MACOSX"
        ]

        for filename in files:
            if filename.startswith("._"):
                continue

            extension = os.path.splitext(filename)[1].lower()

            if extension in required:
                found.setdefault(extension, []).append(
                    os.path.join(root, filename)
                )

    missing = sorted(required - set(found))

    return {
        "found": found,
        "missing": missing,
    }

# =========================================================
# EXCEL SHEET
# =========================================================

def get_metadata_sheet(
    workbook,
):
    """
    Find the metadata worksheet.
    """

    preferred_names = [
        "Meatadata_State_2001",
        "Metadata_State_2001",
        "Metadata",
    ]

    for sheet_name in preferred_names:

        if sheet_name in workbook.sheetnames:

            return workbook[sheet_name]

    return workbook[workbook.sheetnames[0]]


# =========================================================
# TEXT NORMALIZATION
# =========================================================

def normalize_text(
    value: Any,
) -> str:
    """
    Normalize Excel labels for comparison.
    """

    if value is None:

        return ""

    return (
        str(value)
        .strip()
        .lower()
        .replace("_", " ")
        .replace("-", " ")
        .replace("/", " ")
    )


# =========================================================
# HISTORY FIELD HELPERS
# =========================================================

def get_main_field_value(
    main_fields: List[Dict[str, Any]],
    field_name: str,
) -> str:
    """
    Return the exact value of a main metadata field.

    This is intentionally used for History so values such as:

        Entity Label     -> State
        Publication Date -> 31st March 2025 (Current Date)

    are stored exactly as they appear in the reviewed metadata.
    """

    target = normalize_text(field_name)

    if not target:
        return ""

    for field in main_fields:
        current_name = normalize_text(
            field.get("name", "")
        )

        if current_name == target:
            value = field.get("value", "")

            if value is None:
                return ""

            return str(value).strip()

    return ""


# =========================================================
# MERGED CELL HELPERS
# =========================================================

def get_merged_range(
    worksheet,
    row: int,
    column: int,
):
    """
    Return merged range containing a cell.
    """

    coordinate = worksheet.cell(
        row=row,
        column=column,
    ).coordinate

    for merged_range in worksheet.merged_cells.ranges:

        if coordinate in merged_range:

            return merged_range

    return None


def is_writable_cell(
    worksheet,
    row: int,
    column: int,
) -> bool:
    """
    Check whether an Excel cell can be written.

    The check uses row/column numbers of the merged range,
    not its stored start-cell object. After insert_rows()
    openpyxl moves cells but not merged ranges, so the
    stored start cell can point to the wrong row.
    """

    cell = worksheet.cell(
        row=row,
        column=column,
    )

    if isinstance(
        cell,
        MergedCell,
    ):

        return False

    merged_range = get_merged_range(
        worksheet,
        row,
        column,
    )

    if merged_range:

        return (
            row == merged_range.min_row
            and column == merged_range.min_col
        )

    return True


def write_cell(
    worksheet,
    row: int,
    column: int,
    value: Any,
) -> bool:
    """
    Safely write to Excel.
    """

    if not is_writable_cell(
        worksheet,
        row,
        column,
    ):

        return False

    worksheet.cell(
        row=row,
        column=column,
    ).value = (
        "" if value is None else value
    )

    return True


# =========================================================
# FIND MAIN FIELD ROW
# =========================================================

def find_template_field_row(
    worksheet,
    field_name: str,
) -> Optional[int]:

    target = normalize_text(
        field_name
    )

    if not target:

        return None

    for row in range(
        1,
        worksheet.max_row + 1,
    ):

        value = worksheet.cell(
            row=row,
            column=1,
        ).value

        if normalize_text(
            value
        ) == target:

            return row

    return None


# =========================================================
# UPDATE MAIN FIELDS
# =========================================================

def update_main_fields(
    worksheet,
    main_fields: List[Dict[str, Any]],
):

    updated = 0

    for field in main_fields:

        field_name = str(
            field.get(
                "name",
                "",
            )
        ).strip()

        if not field_name:

            continue

        value = field.get(
            "value",
            "",
        )

        # -------------------------------------------------
        # Prefer stored template row
        # -------------------------------------------------

        target_row = field.get(
            "template_row"
        )

        try:

            target_row = (
                int(target_row)
                if target_row
                else None
            )

        except (
            TypeError,
            ValueError,
        ):

            target_row = None

        # -------------------------------------------------
        # Otherwise find by Column A
        # -------------------------------------------------

        if not target_row:

            target_row = find_template_field_row(
                worksheet,
                field_name,
            )

        if not target_row:

            print(
                "WARNING: Main field not found:",
                field_name,
            )

            continue

        # -------------------------------------------------
        # Write value into Column B
        # -------------------------------------------------

        if write_cell(
            worksheet,
            target_row,
            2,
            value,
        ):

            # The template formats some value cells as "0"
            # (whole numbers), which would display
            # 68.11008984 as 68. Decimal numbers (such as the
            # bounding coordinates) are given the General
            # format so the full value is shown.
            if (
                isinstance(value, float)
                and not value.is_integer()
            ):

                worksheet.cell(
                    row=target_row,
                    column=2,
                ).number_format = "General"

            updated += 1

    print(
        f"Main metadata fields updated: "
        f"{updated}/{len(main_fields)}"
    )


# =========================================================
# FIND ATTRIBUTE HEADER ROWS
# =========================================================

def find_attribute_header_rows(
    worksheet,
) -> List[int]:
    """
    Find every existing Attribute #N row.
    """

    rows = []

    for row in range(
        1,
        worksheet.max_row + 1,
    ):

        value = worksheet.cell(
            row=row,
            column=1,
        ).value

        text = normalize_text(
            value
        )

        if text.startswith(
            "attribute #"
        ):

            rows.append(row)

    return rows


# =========================================================
# FIND ATTRIBUTE BLOCK
# =========================================================

def get_attribute_block_rows(
    worksheet,
    attribute_index: int,
):
    """
    Return the exact rows belonging to one attribute.
    """

    header_rows = find_attribute_header_rows(
        worksheet
    )

    position = attribute_index - 1

    if position < 0:

        return None

    if position >= len(header_rows):

        return None

    start_row = header_rows[position]

    if position + 1 < len(header_rows):

        end_row = (
            header_rows[position + 1] - 1
        )

    else:

        end_row = worksheet.max_row

    return (
        start_row,
        end_row,
    )

# =========================================================
# REMOVE ALL BACKGROUND COLORS
# =========================================================

def remove_all_background_colors(
    worksheet,
):
    """
    Remove all background/fill colors from the
    final generated Excel sheet.

    This keeps:
        - fields
        - values
        - borders
        - fonts
        - alignment
        - row heights
        - column widths

    Only the background fill/color is removed.
    """

    plain_fill = PatternFill(fill_type=None)

    for row in worksheet.iter_rows():

        for cell in row:

            if isinstance(cell, MergedCell):
                continue

            cell.fill = copy.copy(
                plain_fill
            )

    print(
        "All background colors removed from final Excel."
    )
# =========================================================
# FIND LABEL INSIDE ATTRIBUTE BLOCK
# =========================================================

def find_label_row(
    worksheet,
    start_row: int,
    end_row: int,
    label: str,
) -> Optional[int]:

    target = normalize_text(
        label
    )

    for row in range(
        start_row,
        end_row + 1,
    ):

        value = worksheet.cell(
            row=row,
            column=1,
        ).value

        if normalize_text(
            value
        ) == target:

            return row

    return None


# =========================================================
# WRITE ATTRIBUTE VALUE
# =========================================================

def write_attribute_value(
    worksheet,
    start_row: int,
    end_row: int,
    label: str,
    value: Any,
):
    """
    Find label in Column A and write its
    corresponding value to Column B.
    """

    row = find_label_row(
        worksheet,
        start_row,
        end_row,
        label,
    )

    if row is None:

        print(
            f"WARNING: Label not found: {label}"
        )

        return False

    return write_cell(
        worksheet,
        row,
        2,
        value,
    )


# =========================================================
# WRITE ATTRIBUTE SOURCE
# =========================================================

def write_attribute_source(
    worksheet,
    start_row: int,
    end_row: int,
    label: str,
    source: str,
):
    """
    Write internal source information into Column C.

    Column C is temporary and will be completely removed
    from the final generated Excel.
    """

    row = find_label_row(
        worksheet,
        start_row,
        end_row,
        label,
    )

    if row is None:

        return False

    return write_cell(
        worksheet,
        row,
        3,
        source,
    )


# =========================================================
# COPY CELL STYLE
# =========================================================

def copy_cell_style(
    source_cell,
    target_cell,
):
    """
    Copy all important formatting from one
    Excel cell to another.
    """

    if isinstance(
        source_cell,
        MergedCell,
    ):

        return

    target_cell._style = copy.copy(
        source_cell._style
    )

    if source_cell.has_style:

        target_cell.font = copy.copy(
            source_cell.font
        )

        target_cell.fill = copy.copy(
            source_cell.fill
        )

        target_cell.border = copy.copy(
            source_cell.border
        )

        target_cell.alignment = copy.copy(
            source_cell.alignment
        )

        target_cell.protection = copy.copy(
            source_cell.protection
        )

    target_cell.number_format = (
        source_cell.number_format
    )

    if source_cell.hyperlink:

        target_cell._hyperlink = copy.copy(
            source_cell.hyperlink
        )

    if source_cell.comment:

        target_cell.comment = copy.copy(
            source_cell.comment
        )


# =========================================================
# COPY ROW DIMENSION
# =========================================================

def copy_row_dimension(
    worksheet,
    source_row: int,
    target_row: int,
):
    """
    Copy row height / hidden / outline
    information.
    """

    source_dimension = (
        worksheet.row_dimensions[
            source_row
        ]
    )

    target_dimension = (
        worksheet.row_dimensions[
            target_row
        ]
    )

    if source_dimension.height is not None:

        target_dimension.height = (
            source_dimension.height
        )

    target_dimension.hidden = (
        source_dimension.hidden
    )

    target_dimension.outlineLevel = (
        source_dimension.outlineLevel
    )

    target_dimension.collapsed = (
        source_dimension.collapsed
    )


# =========================================================
# COPY ATTRIBUTE BLOCK
# =========================================================

def copy_attribute_block(
    worksheet,
    source_start_row: int,
    source_end_row: int,
    target_start_row: int,
):
    """
    Copy an existing attribute block and insert it
    as a new attribute block.
    """

    block_height = (
        source_end_row
        - source_start_row
        + 1
    )

    # -----------------------------------------------------
    # Insert rows
    # -----------------------------------------------------

    worksheet.insert_rows(
        target_start_row,
        amount=block_height,
    )

    # -----------------------------------------------------
    # Copy cell contents and styles
    # -----------------------------------------------------

    max_column = worksheet.max_column

    for offset in range(
        block_height
    ):

        source_row = (
            source_start_row
            + offset
        )

        target_row = (
            target_start_row
            + offset
        )

        copy_row_dimension(
            worksheet,
            source_row,
            target_row,
        )

        for column in range(
            1,
            max_column + 1,
        ):

            source_cell = worksheet.cell(
                row=source_row,
                column=column,
            )

            target_cell = worksheet.cell(
                row=target_row,
                column=column,
            )

            if isinstance(
                source_cell,
                MergedCell,
            ):

                continue

            target_cell.value = (
                source_cell.value
            )

            copy_cell_style(
                source_cell,
                target_cell,
            )

    # -----------------------------------------------------
    # Re-create merged cells
    # -----------------------------------------------------

    source_merges = []

    for merged_range in list(
        worksheet.merged_cells.ranges
    ):

        min_row = merged_range.min_row
        max_row = merged_range.max_row

        if (
            min_row >= source_start_row
            and max_row <= source_end_row
        ):

            source_merges.append(
                merged_range
            )

    row_shift = (
        target_start_row
        - source_start_row
    )

    for merged_range in source_merges:

        new_min_row = (
            merged_range.min_row
            + row_shift
        )

        new_max_row = (
            merged_range.max_row
            + row_shift
        )

        new_min_col = (
            merged_range.min_col
        )

        new_max_col = (
            merged_range.max_col
        )

        start_coordinate = (
            f"{get_column_letter(new_min_col)}"
            f"{new_min_row}"
        )

        end_coordinate = (
            f"{get_column_letter(new_max_col)}"
            f"{new_max_row}"
        )

        new_range = (
            f"{start_coordinate}:{end_coordinate}"
        )

        already_exists = any(
            str(existing_range)
            == new_range
            for existing_range
            in worksheet.merged_cells.ranges
        )

        if not already_exists:

            worksheet.merge_cells(
                new_range
            )

    # -----------------------------------------------------
    # Clear copied attribute values
    #
    # Keep Column A labels.
    # Clear B/C because these belong to the
    # new attribute.
    # -----------------------------------------------------

    for row in range(
        target_start_row,
        target_start_row + block_height,
    ):

        if is_writable_cell(
            worksheet,
            row,
            2,
        ):

            worksheet.cell(
                row=row,
                column=2,
            ).value = ""

        if is_writable_cell(
            worksheet,
            row,
            3,
        ):

            worksheet.cell(
                row=row,
                column=3,
            ).value = ""

    return (
        target_start_row,
        target_start_row
        + block_height
        - 1,
    )


# =========================================================
# EXPAND ATTRIBUTE TEMPLATE
# =========================================================

def ensure_attribute_capacity(
    worksheet,
    required_attributes: int,
):
    """
    Ensure that the worksheet contains enough
    Attribute # blocks.
    """

    if required_attributes <= 0:

        return

    header_rows = (
        find_attribute_header_rows(
            worksheet
        )
    )

    if not header_rows:

        raise ValueError(
            "No Attribute # rows were found "
            "in the metadata template."
        )

    current_count = len(
        header_rows
    )

    print(
        "Template currently contains "
        f"{current_count} attribute blocks."
    )

    print(
        "Required attribute blocks: "
        f"{required_attributes}"
    )

    if (
        current_count
        >= required_attributes
    ):

        print(
            "Template already has enough "
            "attribute blocks."
        )

        return

    # -----------------------------------------------------
    # Remove merged ranges in the attribute area.
    #
    # openpyxl does not move merged ranges when rows are
    # inserted, so they would end up on the wrong rows and
    # block the "Attribute #N" headers from being renumbered
    # (e.g. #3, #3, #5).
    # -----------------------------------------------------

    first_attribute_row = header_rows[0]

    for merged_range in list(
        worksheet.merged_cells.ranges
    ):

        if merged_range.max_row >= first_attribute_row:

            worksheet.unmerge_cells(
                str(merged_range)
            )

    # -----------------------------------------------------
    # Use LAST existing block as source
    # -----------------------------------------------------

    source_start_row = (
        header_rows[-1]
    )

    if len(header_rows) >= 2:

        previous_start_row = (
            header_rows[-2]
        )

        block_height = (
            source_start_row
            - previous_start_row
        )

        source_end_row = (
            source_start_row
            + block_height
            - 1
        )

    else:

        source_end_row = (
            worksheet.max_row
        )

    insertion_row = (
        source_end_row + 1
    )

    attributes_to_create = (
        required_attributes
        - current_count
    )

    print(
        "Creating "
        f"{attributes_to_create} "
        "additional attribute blocks..."
    )

    for i in range(
        attributes_to_create
    ):

        print(
            "Creating attribute block "
            f"#{current_count + i + 1}"
        )

        new_start, new_end = (
            copy_attribute_block(
                worksheet=worksheet,
                source_start_row=source_start_row,
                source_end_row=source_end_row,
                target_start_row=insertion_row,
            )
        )

        source_start_row = new_start

        source_end_row = new_end

        insertion_row = (
            new_end + 1
        )

    # -----------------------------------------------------
    # Rename all headers
    # -----------------------------------------------------

    final_header_rows = (
        find_attribute_header_rows(
            worksheet
        )
    )

    for index, row in enumerate(
        final_header_rows,
        start=1,
    ):

        write_cell(
            worksheet,
            row,
            1,
            f"Attribute #{index}",
        )

    print(
        "Attribute template expanded successfully."
    )

    print(
        "Total attribute blocks now: "
        f"{len(final_header_rows)}"
    )


# =========================================================
# MISSING FIELD CHECK / RED HEADER
# =========================================================

RED_FONT_COLOR = "FFFF0000"


def is_blank(
    value: Any,
) -> bool:
    """
    True when a value is missing.

    0 is a real value (for example a range minimum),
    so only None and empty text count as blank.
    """

    if value is None:

        return True

    return str(value).strip() == ""


def get_missing_attribute_fields(
    attribute: Dict[str, Any],
) -> List[str]:
    """
    Return the names of the attribute fields that are empty.

    Checked for every attribute:
        Attribute Label, Attribute Definition,
        Attribute Definition Source,
        Attribute Unit of Measurement

    Range Domain Minimum / Maximum are checked only for
    numeric attributes (unit = "Numbers"). Text columns
    have no numeric range, so a blank range is normal.
    """

    missing: List[str] = []

    if is_blank(attribute.get("name")):
        missing.append("Attribute Label")

    if is_blank(attribute.get("definition")):
        missing.append("Attribute Definition")

    if is_blank(attribute.get("definition_source")):
        missing.append("Attribute Definition Source")

    if is_blank(attribute.get("unit")):
        missing.append("Attribute Unit of Measurement")

    if str(attribute.get("unit", "")).strip().lower() == "numbers":

        if is_blank(attribute.get("range_min")):
            missing.append("Range Domain Minimum")

        if is_blank(attribute.get("range_max")):
            missing.append("Range Domain Maximum")

    return missing


def mark_attribute_header_red(
    worksheet,
    header_row: int,
):
    """
    Make the "Attribute #N" text in Column A red,
    keeping the cell's existing font (name, size, bold).
    """

    cell = worksheet.cell(
        row=header_row,
        column=1,
    )

    if isinstance(
        cell,
        MergedCell,
    ):

        return

    red_font = copy.copy(
        cell.font
    )

    red_font.color = Color(
        rgb=RED_FONT_COLOR
    )

    cell.font = red_font


# =========================================================
# WRITE ONE ATTRIBUTE
# =========================================================

def write_single_attribute(
    worksheet,
    attribute_number: int,
    attribute: Dict[str, Any],
):
    """
    Write one complete attribute into its
    corresponding template block.
    """

    block = get_attribute_block_rows(
        worksheet,
        attribute_number,
    )

    if block is None:

        print(
            f"WARNING: Attribute #{attribute_number} "
            "does not exist in template."
        )

        return False

    start_row, end_row = block

    name = attribute.get(
        "name",
        "",
    )

    definition = attribute.get(
        "definition",
        "",
    )

    definition_source = attribute.get(
        "definition_source",
        "",
    )

    range_min = attribute.get(
        "range_min",
        "",
    )

    range_max = attribute.get(
        "range_max",
        "",
    )

    unit = attribute.get(
        "unit",
        "",
    )

    # =====================================================
    # ATTRIBUTE HEADER
    # =====================================================

    header_written = write_cell(
        worksheet,
        start_row,
        1,
        f"Attribute #{attribute_number}",
    )

    # =====================================================
    # ATTRIBUTE LABEL
    # =====================================================

    write_attribute_value(
        worksheet,
        start_row,
        end_row,
        "Attribute Label",
        name,
    )

    write_attribute_source(
        worksheet,
        start_row,
        end_row,
        "Attribute Label",
        "From Shape File",
    )

    # =====================================================
    # DEFINITION
    # =====================================================

    write_attribute_value(
        worksheet,
        start_row,
        end_row,
        "Attribute Definition",
        definition,
    )

    write_attribute_source(
        worksheet,
        start_row,
        end_row,
        "Attribute Definition",
        (
            "From Master"
            if definition
            else ""
        ),
    )

    # =====================================================
    # DEFINITION SOURCE
    # =====================================================

    write_attribute_value(
        worksheet,
        start_row,
        end_row,
        "Attribute Definition Source",
        definition_source,
    )

    write_attribute_source(
        worksheet,
        start_row,
        end_row,
        "Attribute Definition Source",
        (
            "From Master"
            if definition_source
            else ""
        ),
    )

    # =====================================================
    # RANGE MINIMUM
    # =====================================================

    write_attribute_value(
        worksheet,
        start_row,
        end_row,
        "Range Domain Minimum",
        range_min,
    )

    write_attribute_source(
        worksheet,
        start_row,
        end_row,
        "Range Domain Minimum",
        "From Shape File",
    )

    # =====================================================
    # RANGE MAXIMUM
    # =====================================================

    write_attribute_value(
        worksheet,
        start_row,
        end_row,
        "Range Domain Maximum",
        range_max,
    )

    write_attribute_source(
        worksheet,
        start_row,
        end_row,
        "Range Domain Maximum",
        "From Shape File",
    )

    # =====================================================
    # UNIT
    # =====================================================

    write_attribute_value(
        worksheet,
        start_row,
        end_row,
        "Attribute Unit of Measurement",
        unit,
    )

    write_attribute_source(
        worksheet,
        start_row,
        end_row,
        "Attribute Unit of Measurement",
        "From Shape File",
    )

    # =====================================================
    # MARK INCOMPLETE ATTRIBUTES IN RED
    #
    # If any required field of this attribute is missing,
    # the "Attribute #N" header text is shown in red.
    # =====================================================

    missing_fields = get_missing_attribute_fields(
        attribute
    )

    if missing_fields:

        mark_attribute_header_red(
            worksheet,
            start_row,
        )

    print(
        f"Attribute #{attribute_number}: "
        f"{name} -> rows "
        f"{start_row}-{end_row}"
        + (
            f" | MISSING: {', '.join(missing_fields)}"
            if missing_fields
            else ""
        )
    )

    return header_written


# =========================================================
# REMOVE UNUSED ATTRIBUTE BLOCKS
# =========================================================

def remove_unused_attribute_blocks(
    worksheet,
    used_attributes: int,
):
    """
    Keep only the attribute blocks that belong to columns
    present in the shapefile.

    The template ships with sample attribute blocks. If the
    shapefile has fewer columns than the template has blocks,
    the surplus sample blocks (with old template data) are
    deleted.
    """

    header_rows = find_attribute_header_rows(
        worksheet
    )

    if len(header_rows) <= used_attributes:

        return

    first_unused_row = header_rows[
        used_attributes
    ]

    # Merged ranges do not move with deleted rows,
    # so remove any in the area being deleted.
    for merged_range in list(
        worksheet.merged_cells.ranges
    ):

        if merged_range.max_row >= first_unused_row:

            worksheet.unmerge_cells(
                str(merged_range)
            )

    rows_to_delete = (
        worksheet.max_row
        - first_unused_row
        + 1
    )

    worksheet.delete_rows(
        first_unused_row,
        rows_to_delete,
    )

    print(
        "Removed unused template attribute blocks: "
        f"{len(header_rows) - used_attributes} "
        f"(rows {first_unused_row} to end)"
    )


# =========================================================
# REMOVE EMPTY ROWS AT THE END OF THE SHEET
# =========================================================

def remove_trailing_blank_rows(
    worksheet,
):
    """
    Delete rows at the bottom of the sheet where Column A and
    Column B are both empty (for example the template's closing
    bordered row).
    """

    removed = 0

    while worksheet.max_row > 1:

        last_row = worksheet.max_row

        value_a = worksheet.cell(
            row=last_row,
            column=1,
        ).value

        value_b = worksheet.cell(
            row=last_row,
            column=2,
        ).value

        if (
            str(value_a or "").strip()
            or str(value_b or "").strip()
        ):

            break

        # A merged range on this row would be left behind
        for merged_range in list(
            worksheet.merged_cells.ranges
        ):

            if (
                merged_range.min_row <= last_row
                <= merged_range.max_row
            ):

                worksheet.unmerge_cells(
                    str(merged_range)
                )

        worksheet.delete_rows(
            last_row,
            1,
        )

        removed += 1

    if removed:

        print(
            f"Removed {removed} empty row(s) at the end."
        )


# =========================================================
# WRITE ALL ATTRIBUTES
# =========================================================

def write_attributes(
    worksheet,
    attributes: List[Dict[str, Any]],
):
    """
    Write all shapefile attributes.
    """

    if not attributes:

        print(
            "No shapefile attributes to write."
        )

        return

    ensure_attribute_capacity(
        worksheet,
        len(attributes),
    )

    template_attribute_rows = (
        find_attribute_header_rows(
            worksheet
        )
    )

    if (
        len(template_attribute_rows)
        < len(attributes)
    ):

        raise ValueError(
            "Unable to create enough "
            "attribute sections. "
            f"Required: {len(attributes)}, "
            f"available: "
            f"{len(template_attribute_rows)}."
        )

    print(
        "========================================"
    )

    print(
        "FINAL ATTRIBUTE TEMPLATE ROWS:"
    )

    print(
        template_attribute_rows
    )

    print(
        "========================================"
    )

    written = 0

    for index, attribute in enumerate(
        attributes,
        start=1,
    ):

        if write_single_attribute(
            worksheet,
            index,
            attribute,
        ):

            written += 1

    print(
        f"Attributes written: "
        f"{written}/{len(attributes)}"
    )

    # Keep only the attributes present in the shapefile
    remove_unused_attribute_blocks(
        worksheet,
        len(attributes),
    )

    # The template ends with an empty bordered row after the
    # last attribute block. It gets pushed to the bottom when
    # blocks are added, so remove it from the final file.
    remove_trailing_blank_rows(
        worksheet
    )


# =========================================================
# REMOVE ALL COLUMNS AFTER B
# =========================================================

def remove_extra_columns(
    worksheet,
):
    """
    Remove every column after Column B.

    Final generated Excel will contain ONLY:

        Column A = FIELD
        Column B = VALUE

    Columns C, D, E, F, etc. are removed completely.

    The template Remarks column and any other helper
    columns are therefore not present in the final file.
    """

    keep_columns = 2

    current_columns = (
        worksheet.max_column
    )

    if current_columns <= keep_columns:

        print(
            "No extra columns to remove."
        )

        return

    columns_to_remove = (
        current_columns
        - keep_columns
    )

    print(
        "Removing extra columns:"
    )

    print(
        f"Current columns: {current_columns}"
    )

    print(
        f"Keeping columns: A and B"
    )

    print(
        f"Removing columns: C through "
        f"{get_column_letter(current_columns)}"
    )

    worksheet.delete_cols(
        keep_columns + 1,
        columns_to_remove,
    )

    print(
        "Final Excel now contains "
        "only Column A and Column B."
    )


# =========================================================
# UPLOAD SHAPEFILE
# =========================================================

@router.post("/upload")
async def upload_shapefile(
    file: UploadFile = File(...),
):

    if not file.filename:

        raise HTTPException(
            status_code=400,
            detail="No file selected.",
        )

    if not file.filename.lower().endswith(
        ".zip"
    ):

        raise HTTPException(
            status_code=400,
            detail="Please upload a ZIP file.",
        )

    if not os.path.exists(
        TEMPLATE_PATH
    ):

        raise HTTPException(
            status_code=500,
            detail=(
                "Metadata template not found: "
                f"{TEMPLATE_PATH}"
            ),
        )

    upload_id = (
        next(
            tempfile._get_candidate_names()
        )
    )

    work_dir = os.path.join(
        UPLOAD_ROOT,
        upload_id,
    )

    os.makedirs(
        work_dir,
        exist_ok=True,
    )

    zip_path = os.path.join(
        work_dir,
        os.path.basename(file.filename),
    )

    extracted_dir = os.path.join(
        work_dir,
        "extracted",
    )

    os.makedirs(
        extracted_dir,
        exist_ok=True,
    )

    try:

        # =================================================
        # SAVE ZIP
        # =================================================

        with open(
            zip_path,
            "wb",
        ) as output:

            shutil.copyfileobj(
                file.file,
                output,
            )

        # =================================================
        # VALIDATE ZIP
        # =================================================

        try:

            with zipfile.ZipFile(
                zip_path,
                "r",
            ) as zip_ref:

                bad_file = (
                    zip_ref.testzip()
                )

                if bad_file:

                    raise HTTPException(
                        status_code=400,
                        detail=(
                            "The uploaded ZIP "
                            "is corrupted."
                        ),
                    )

                # Block "zip slip" (entries such as ../../x)
                safe_base = os.path.realpath(extracted_dir)

                for member in zip_ref.namelist():

                    member_target = os.path.realpath(
                        os.path.join(extracted_dir, member)
                    )

                    if (
                        member_target != safe_base
                        and not member_target.startswith(
                            safe_base + os.sep
                        )
                    ):

                        raise HTTPException(
                            status_code=400,
                            detail=(
                                "The ZIP contains an unsafe file path."
                            ),
                        )

                zip_ref.extractall(
                    extracted_dir
                )

        except zipfile.BadZipFile:

            raise HTTPException(
                status_code=400,
                detail=(
                    "The uploaded file "
                    "is not a valid ZIP file."
                ),
            )

        
        # =================================================
        # VALIDATE ALL FOUR SHAPEFILE COMPONENTS
        # =================================================

        components = find_shapefile_components(
            extracted_dir
        )

        missing_files = components["missing"]

        if missing_files:
            missing_text = ", ".join(missing_files)

            raise HTTPException(
                status_code=400,
                detail=(
                    f"Missing required shapefile file(s): "
                    f"{missing_text}. The ZIP must contain "
                    ".shp, .shx, .dbf, and .prj files."
                ),
            )

        shp_path = components["found"][".shp"][0]
        dbf_path = components["found"][".dbf"][0]

        print(f"SHP found: {shp_path}")
        print(f"SHX found: {components['found']['.shx'][0]}")
        print(f"DBF found: {dbf_path}")
        print(f"PRJ found: {components['found']['.prj'][0]}")

        print(f"Template found: {TEMPLATE_PATH}")

     
        # =================================================
        # BUILD METADATA
        # =================================================

        metadata = build_metadata(
            dbf_path=dbf_path,
            shp_path=shp_path,
            master_path=TEMPLATE_PATH,
        )

        # =================================================
        # INTERNAL INFORMATION
        # =================================================

        metadata["_template_path"] = (
            TEMPLATE_PATH
        )

        metadata["_upload_id"] = (
            upload_id
        )

        metadata["filename"] = (
            os.path.splitext(
                file.filename
            )[0]
        )

        # =================================================
        # DEBUG
        # =================================================

        print(
            "========================================"
        )

        print(
            "MAIN FIELDS:",
            len(
                metadata.get(
                    "main_fields",
                    [],
                )
            ),
        )

        print(
            "ATTRIBUTES:",
            len(
                metadata.get(
                    "fields",
                    [],
                )
            ),
        )

        print(
            "RECORDS:",
            metadata.get(
                "total_records",
                0,
            ),
        )

        print(
            "========================================"
        )

        return metadata

    except HTTPException:

        raise

    except Exception as exc:

        print(
            "Metadata generation error:",
            repr(exc),
        )

        raise HTTPException(
            status_code=500,
            detail=str(exc),
        )

    finally:

        # The ZIP and the extracted shapefile are only needed
        # while the metadata is being built. Remove them now,
        # whether the upload succeeded or failed.
        shutil.rmtree(
            work_dir,
            ignore_errors=True,
        )


# =========================================================
# GENERATE EXCEL
# =========================================================

@router.post("/generate-excel")
async def generate_metadata_excel(
    metadata: Dict[str, Any],
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not os.path.exists(
        TEMPLATE_PATH
    ):

        raise HTTPException(
            status_code=500,
            detail=(
                "Metadata template not found: "
                f"{TEMPLATE_PATH}"
            ),
        )

    try:

        # =================================================
        # ALWAYS LOAD ORIGINAL TEMPLATE
        # =================================================

        workbook = openpyxl.load_workbook(
            TEMPLATE_PATH
        )

        worksheet = get_metadata_sheet(
            workbook
        )

        print(
            "========================================"
        )

        print(
            "GENERATING METADATA EXCEL"
        )

        print(
            f"Template: {TEMPLATE_PATH}"
        )

        print(
            f"Worksheet: {worksheet.title}"
        )

        print(
            "========================================"
        )

        # =================================================
        # MAIN FIELDS
        # =================================================

        main_fields = metadata.get(
            "main_fields",
            [],
        )

        update_main_fields(
            worksheet,
            main_fields,
        )

        # =================================================
        # ATTRIBUTES
        # =================================================

        attributes = metadata.get(
            "fields",
            [],
        )

        print(
            f"Writing {len(attributes)} "
            "attributes..."
        )

        write_attributes(
            worksheet,
            attributes,
        )

        # =================================================
        # IMPORTANT
        #
        # Column C and every column after B are
        # template/helper columns.
        #
        # Remove them ONLY AFTER all metadata
        # has been processed.
        # =================================================

        remove_extra_columns(
            worksheet
        )

        # =================================================
        # OUTPUT DIRECTORY
        # =================================================

        output_dir = GENERATED_DIR

        os.makedirs(
            output_dir,
            exist_ok=True,
        )

        # =================================================
        # OUTPUT FILENAME
        # =================================================

        filename = metadata.get(
            "filename",
            "GIS_Metadata",
        )

        safe_filename = (
            str(filename)
            .replace(
                "/",
                "_",
            )
            .replace(
                "\\",
                "_",
            )
            .strip()
        )

        if not safe_filename:

            safe_filename = (
                "GIS_Metadata"
            )

        output_filename = (
            f"{safe_filename}"
            "_metadata.xlsx"
        )

        # Unique stored name: two exports with the same name no
        # longer overwrite each other, and deleting one History
        # record can never delete another record's file.
        output_path = os.path.join(
            output_dir,
            f"{safe_filename}_{uuid.uuid4().hex[:8]}_metadata.xlsx",
        )
        # =================================================
# REMOVE ALL BACKGROUND COLORS
# =================================================

        remove_all_background_colors(
         worksheet
        )
        # =================================================
        # SAVE
        # =================================================
      
        workbook.save(
            output_path
        )

        # =================================================
        # SAVE GENERATION HISTORY TO POSTGRESQL
        # =================================================
        #
        # Entity:
        #   Entity Label -> e.g. "State"
        #
        # Year column in History:
        #   Publication Date -> e.g.
        #   "31st March 2025 (Current Date)"
        #
        # The complete publication-date text is preserved.
        # We do NOT extract only the numeric year.
        # =================================================

        entity = get_main_field_value(
            main_fields,
            "Entity Label",
        )

        publication_date = get_main_field_value(
            main_fields,
            "Publication Date",
        )

        history_record = GeneratedFileHistory(
            user_id=current_user.id,
            owner_username=current_user.username,
            filename=output_filename,
            original_filename=metadata.get("filename"),
            entity=entity or None,
            publication_date=publication_date or None,
            records_count=int(
                metadata.get("total_records", 0) or 0
            ),
            attributes_count=len(attributes),
            status="success",
            generated_file_path=output_path,
        )

        try:
            db.add(history_record)
            db.commit()
            db.refresh(history_record)

        except Exception:
            db.rollback()

            # Do not leave an Excel file without its history record.
            if os.path.exists(output_path):
                try:
                    os.remove(output_path)
                except OSError:
                    pass

            raise

        print(
            "========================================"
        )

        print(
            "History saved to PostgreSQL."
        )

        print(
            f"History ID: {history_record.id}"
        )

        print(
            f"Entity: {history_record.entity}"
        )

        print(
            f"Publication Date: "
            f"{history_record.publication_date}"
        )

        print(
            "========================================"
        )

        print(
            "========================================"
        )

        print(
            "Excel generated successfully:"
        )

        print(
            output_path
        )

        print(
            "Final columns: A and B only"
        )

        print(
            "========================================"
        )

        # =================================================
        # RETURN FILE
        # =================================================

        return FileResponse(
            path=output_path,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            filename=output_filename,
        )

    except HTTPException:

        raise

    except Exception as exc:

        print(
            "Excel generation error:",
            repr(exc),
        )

        raise HTTPException(
            status_code=500,
            detail=str(exc),
        )