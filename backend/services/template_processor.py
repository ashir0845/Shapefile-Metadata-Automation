import openpyxl


REMARKS_EDITABLE = "editable"


def normalize(value):
    if value is None:
        return ""

    return str(value).strip()


def is_editable(remarks):
    return normalize(remarks).lower() == REMARKS_EDITABLE


def load_template_metadata(template_path):
    """
    Reads the metadata template and returns all metadata rows.

    Template structure:
        Column A -> FIELD
        Column B -> VALUE
        Column C -> Remarks
    """

    workbook = openpyxl.load_workbook(
        template_path,
        data_only=True
    )

    if "Meatadata_State_2001" not in workbook.sheetnames:
        raise ValueError(
            "Sheet 'Meatadata_State_2001' was not found "
            "in the metadata template."
        )

    worksheet = workbook["Meatadata_State_2001"]

    metadata = []

    for row_number in range(2, worksheet.max_row + 1):

        field = worksheet.cell(
            row_number,
            1
        ).value

        value = worksheet.cell(
            row_number,
            2
        ).value

        remarks = worksheet.cell(
            row_number,
            3
        ).value

        if not field:
            continue

        field = normalize(field)
        remarks = normalize(remarks)

        metadata.append(
            {
                "row_number": row_number,
                "field": field,
                "value": value,
                "remarks": remarks,
                "editable": is_editable(remarks),
            }
        )

    return metadata